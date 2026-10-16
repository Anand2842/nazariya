"""Builds nazariya-unit-economics.xlsx (live formulas) and csv/*.csv (computed values).

Every number on the Inputs sheet is editable (blue). All other sheets are formulas on top of it, so changing an
input (e.g. GPU price, capturer salary) reprices everything when the workbook is opened in Excel / Numbers /
Google Sheets. The CSVs are a snapshot of the same model computed here (a tiny evaluator mirrors the formulas).

Run: python3 tools/build_pricing.py   (from nazariya-business/)
"""
import csv
import math
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent.parent
BLUE = Font(color="1F4E9E")
BOLD = Font(bold=True)
HEAD = PatternFill("solid", fgColor="E8EEF7")

# ---------------------------------------------------------------- model definition
# Inputs: key, label, value, unit, source / confidence
INPUTS = [
    ("General", None, None, None, None),
    ("fx", "Exchange rate", 88, "₹ per US$", "Planning assumption (same as research docs)"),
    ("gst", "GST on our invoices (not included in prices)", 0.18, "rate", "India GST on SaaS / services"),
    ("rooms_home", "Rooms in a typical 2BHK tour", 6, "rooms", "2 bed + living + kitchen + 2 bath (research doc)"),
    ("sqft_home", "Carpet area of a typical 2BHK", 900, "sq ft", "Assumption"),
    ("Processing & hosting (measured on this app)", None, None, None, None),
    ("stitch_cpu_s", "CPU time to stitch one 30-shot room", 15, "seconds", "MEASURED: your room 'Hc' stitched in 14.9 s"),
    ("proc_overhead", "Overhead multiplier (tiles, thumbnails, retries)", 2, "x", "Assumption"),
    ("cpu_usd_hr", "Cloud CPU price", 0.04, "US$ per vCPU-hour", "Assumption: typical on-demand cloud vCPU; verify with provider"),
    ("raw_mb_room", "Raw capture kept per room (30 photos)", 3.9, "MB", "MEASURED: priv/capture/mt7wn6pvq9aj = 3.9 MB"),
    ("depth_mb_room", "Depth points per room (depth mode only)", 6, "MB", "App limit is 8 MB per room"),
    ("pub_mb_room", "Published files per room (panorama + tiles)", 1.0, "MB", "MEASURED: 29 MB for 30 published panoramas"),
    ("raw_retention_m", "Months raw captures are kept", 3, "months", "Policy (DPDP minimisation, PRODUCTION.md)"),
    ("r2_gb_usd", "Object storage (Cloudflare R2 Standard)", 0.015, "US$ per GB-month", "developers.cloudflare.com/r2/pricing (egress free)"),
    ("r2_classb_usd_m", "Storage read requests (R2 Class B)", 0.36, "US$ per million", "developers.cloudflare.com/r2/pricing"),
    ("views_tour_m", "Visitor sessions per tour per month", 300, "sessions", "Assumption: replace with real analytics"),
    ("req_per_view", "File requests per visitor session", 15, "requests", "Assumption (tiles + JSON)"),
    ("3D walkthrough (GPU)", None, None, None, None),
    ("gpu_usd_hr", "GPU rental (RTX 4090, RunPod Secure Cloud)", 0.69, "US$ per hour", "RunPod ~US$0.69 secure / ~0.34 community (third-party 2026 listings)"),
    ("gpu_min_walk", "GPU minutes per home walkthrough", 40, "minutes", "Assumption: NOT measured, recon scripts never run on a GPU yet"),
    ("gpu_retry", "Retry / failed-run allowance", 0.3, "rate", "Assumption"),
    ("walk_dataset_mb", "Walkthrough photos kept (dataset)", 120, "MB", "~400 keyframes x 0.3 MB"),
    ("scene_mb", "Published 3D scene size", 40, "MB", "Assumption (demo splats avg 4.5 MB per room)"),
    ("Done-for-you capture (our team shoots)", None, None, None, None),
    ("capturer_month", "Field capturer cost (salary + conveyance)", 30000, "₹ per month", "Assumption: ₹25k salary + ₹5k travel (research doc)"),
    ("homes_per_day", "Homes one capturer shoots per day", 4, "homes", "Assumption: clustered localities"),
    ("days_month", "Working days per month", 24, "days", ""),
    ("travel_visit", "Local travel per visit", 150, "₹", "Assumption"),
    ("camera_inr", "360 camera (Insta360 X5) for hard rooms", 54990, "₹", "Insta360 X5 India launch price ₹54,990"),
    ("camera_life", "Captures per camera before replacement", 600, "captures", "Assumption"),
    ("qa_min_home", "Ops review time per home (hotspots, floor plan, QA)", 30, "minutes", "Assumption"),
    ("ops_hr_inr", "Ops staff cost", 250, "₹ per hour", "Assumption"),
    ("Selling & payments", None, None, None, None),
    ("pg_fee", "Payment gateway fee", 0.02, "rate", "Assumption: ~2% standard Indian gateway rate; verify"),
    ("sales_comm", "Sales commission on service deals", 0.10, "rate", "Policy proposal"),
    ("otp_inr", "OTP SMS price", 0.20, "₹ per SMS", "MSG91 ~₹0.15-0.25 (third-party 2026 comparisons)"),
    ("otp_per_cust_m", "Logins per customer per month", 15, "OTPs", "Assumption"),
    ("support_starter", "Support cost, Starter customer", 100, "₹ per month", "Assumption"),
    ("support_pro", "Support cost, Pro customer", 400, "₹ per month", "Assumption"),
    ("support_business", "Support cost, Business customer", 1500, "₹ per month", "Assumption"),
    ("support_enterprise", "Support cost, Enterprise customer", 8000, "₹ per month", "Assumption"),
    ("tour_churn", "Share of a plan's tours re-captured each month", 0.2, "rate", "Assumption (listings turn over)"),
    ("Fixed monthly costs (lean start)", None, None, None, None),
    ("fx_infra", "Cloud: database, app, worker, storage, monitoring", 25000, "₹ per month", "Research doc range ₹15k-40k"),
    ("fx_tools", "SaaS tools, domain, email", 5000, "₹ per month", "Assumption"),
    ("fx_ca", "Accountant / compliance (CA, GST filing)", 6000, "₹ per month", "Assumption"),
    ("fx_marketing", "Marketing (ads, demo flats, travel)", 25000, "₹ per month", "Assumption"),
    ("start_cash", "Starting cash (founder money)", 300000, "₹", "Assumption: change to your real number"),
    ("vol_scale", "Sales volume scale for Plan12M (1 = plan, 0.5 = downside)", 1, "x", "Scenario switch"),
]

# Unit costs: key, label, formula, unit, how
UNIT = [
    ("cpu_room", "Processing one room (CPU)", "={stitch_cpu_s}*{proc_overhead}/3600*{cpu_usd_hr}*{fx}", "₹ per room", "CPU seconds x price"),
    ("store_raw_room", "Keeping raw photos + depth for the retention period", "=({raw_mb_room}+{depth_mb_room})/1024*{r2_gb_usd}*{raw_retention_m}*{fx}", "₹ per room", "MB -> GB x months"),
    ("store_pub_room_m", "Hosting one published room for a month", "={pub_mb_room}/1024*{r2_gb_usd}*{fx}", "₹ per room-month", ""),
    ("deliv_tour_m", "Delivering a tour to visitors for a month", "={views_tour_m}*{req_per_view}/1000000*{r2_classb_usd_m}*{fx}", "₹ per tour-month", "Egress is free on R2"),
    ("cost_room_360", "One 360° room, self-captured, hosted 12 months", "={cpu_room}+{store_raw_room}+12*{store_pub_room_m}", "₹", ""),
    ("cost_home_diy", "Home tour, self-captured (6 rooms), 12 months", "={rooms_home}*{cost_room_360}+12*{deliv_tour_m}", "₹", "Software cost is close to zero"),
    ("field_home", "Capturer time per home", "={capturer_month}/({homes_per_day}*{days_month})", "₹ per home", ""),
    ("camera_home", "Camera wear per home", "={camera_inr}/{camera_life}", "₹ per home", ""),
    ("qa_home", "Ops review per home", "={qa_min_home}/60*{ops_hr_inr}", "₹ per home", ""),
    ("cost_home_dfy", "Home tour, done-for-you (we shoot), 12 months", "={cost_home_diy}+{field_home}+{travel_visit}+{camera_home}+{qa_home}", "₹", "People cost dominates"),
    ("gpu_walk", "GPU training for one 3D walkthrough", "={gpu_min_walk}/60*{gpu_usd_hr}*{fx}*(1+{gpu_retry})", "₹ per home", "Not yet measured"),
    ("store_walk", "Storing walkthrough dataset + 3D scene", "=({walk_dataset_mb}*{raw_retention_m}+{scene_mb}*12)/1024*{r2_gb_usd}*{fx}", "₹ per home", ""),
    ("cost_walk", "3D walkthrough, total direct cost", "={gpu_walk}+{store_walk}", "₹ per home", ""),
    ("host_tour_m", "Hosting one live home tour for a month", "={rooms_home}*{store_pub_room_m}+{deliv_tour_m}", "₹ per tour-month", ""),
    ("otp_cust_m", "Login SMS per customer per month", "={otp_inr}*{otp_per_cust_m}", "₹ per customer-month", ""),
]

# Price list: key, product, unit, price, direct-cost formula, commission?, per-sqft?, notes
PRICES = [
    ("p_diy", "DIY home tour (self-capture in app, ≤8 rooms, floor plan, embed, 12 months)", "per home / year", 999,
     "={cost_home_diy}", False, True, "Brokers, owners, PG/co-living operators"),
    ("p_room", "Extra room on a DIY tour", "per room", 149, "={cost_room_360}", False, False, ""),
    ("p_dfy", "Done-for-you home tour (we shoot, up to 3BHK)", "per home", 4999,
     "={cost_home_dfy}", True, True, "Fast cash: brokers, builders' resale, rentals"),
    ("p_villa", "Done-for-you large home / villa (≤15 rooms)", "per home", 9999,
     "=2*{field_home}+{travel_visit}+{camera_home}+2*{qa_home}+15*{cost_room_360}+12*{deliv_tour_m}", True, False, "Half-day shoot"),
    ("p_walk", "3D walkthrough add-on (true 3D, GPU)", "per home", 2999,
     "={cost_walk}+{qa_home}", False, False, "Only after GPU pipeline is validated"),
    ("p_dev", "Developer project package (≤4 unit types + amenities, white label, leads, 12 months)", "per project / year", 49999,
     "=5*{cost_home_dfy}+2*{cost_walk}+12*{support_business}", True, False, "Sample flats + clubhouse; replaces site visits"),
    ("p_hotel", "Hotel / resort package (≤10 room categories + lobby, 12 months)", "per property / year", 24999,
     "=2*{cost_home_dfy}+12*{support_pro}", True, False, "OTA + website embed"),
    ("p_renew", "Hosting renewal, any tour", "per tour / year", 499, "=12*{host_tour_m}", False, False, "Year 2 onwards"),
]

# SaaS plans: key, plan, price/month, active tours, seats, walkthroughs incl., support key, features, anchor
PLANS = [
    ("pl_starter", "Starter", 1499, 10, 1, 0, "support_starter", "Self-capture, embeds, Nazariya branding", "TeliportMe Essential US$19-25; Kuula Pro US$20"),
    ("pl_pro", "Pro", 4999, 50, 3, 1, "support_pro", "White label, floor plans, depth mode, leads + analytics", "TeliportMe Business US$29-35; Matterport Pro ~US$55-69"),
    ("pl_business", "Business", 12999, 200, 10, 5, "support_business", "API, custom domain, 5 x 3D walkthroughs/month", "TeliportMe Unlimited US$83-99; Matterport Business ~US$269-355"),
    ("pl_enterprise", "Enterprise (from)", 49999, 1000, 50, 20, "support_enterprise", "SSO, SLA, dedicated success manager", "Custom (all competitors)"),
]

COMPETITORS = [
    # name, plan, usd/month (low), usd/month (high), what you get, source, confidence
    ("TeliportMe", "Essential", 19, 25, "Unlimited tours, 1 user", "https://teliportme.com/pricing (fetched 2026-10-17)", "Reported (two price sets shown; likely annual vs monthly)"),
    ("TeliportMe", "Business", 29, 35, "Unlimited tours, 3 users, 1 custom domain", "https://teliportme.com/pricing", "Reported"),
    ("TeliportMe", "Unlimited", 83, 99, "Unlimited users, 5 custom domains", "https://teliportme.com/pricing", "Reported"),
    ("Kuula", "Pro", 20, 20, "Unlimited uploads, tour editor", "https://www.capterra.com/p/179213/Kuula/pricing/", "Reported (third-party)"),
    ("Kuula", "Business", 36, 48, "Analytics, password tours, custom domain", "Capterra / TrustRadius (conflict)", "Reported (third-party, conflicting)"),
    ("Matterport", "Starter", 12, 14, "5-20 active spaces", "https://travvir.com/blog/en/matterport-pricing-explained", "Reported (third-party); Matterport plans page shows tiers only"),
    ("Matterport", "Professional", 55, 69, "20-150 active spaces", "Third-party listings (conflict)", "Reported; possible Oct-2026 repricing unverified"),
    ("Matterport", "Business", 269, 355, "100-300 active spaces", "Third-party listings (conflict)", "Reported"),
]
COMP_OTHER = [
    ("Nucleus4D", "Per listing", "US$0.03 per sq ft per listing per year + dedicated capture rig", "Podcast summary (research ledger E11)", "Low confidence"),
    ("Indian 360 tour agency", "Per tour", "₹1,00,000 per tour (one Ahmedabad vendor listing)", "https://m.indiamart.com/impcat/360-degree-virtual-tours.html", "Single asking price, not a market rate"),
    ("US 3D tour service", "Per small home", "US$100-500", "https://www.styldod.com/blog/how-much-do-3d-virtual-tours-cost", "Reported"),
]

MONTHS = 12
# 12-month plan volumes (targets you set; blue).
VOLUMES = [
    ("v_dfy", "Done-for-you home tours sold", "p_dfy", [5, 10, 15, 20, 25, 30, 35, 40, 45, 50, 55, 60]),
    ("v_diy", "DIY home tours sold", "p_diy", [0, 5, 10, 20, 30, 45, 60, 75, 90, 110, 130, 150]),
    ("v_villa", "Villas / large homes", "p_villa", [0, 1, 1, 2, 2, 3, 3, 3, 4, 4, 4, 5]),
    ("v_walk", "3D walkthrough add-ons", "p_walk", [0, 0, 0, 1, 2, 3, 4, 5, 6, 7, 8, 9]),
    ("v_dev", "Developer project packages (new)", "p_dev", [0, 0, 1, 1, 1, 2, 2, 2, 2, 3, 3, 3]),
    ("v_hotel", "Hotel packages (new)", "p_hotel", [0, 1, 1, 1, 2, 2, 2, 2, 3, 3, 3, 3]),
    ("v_starter", "Starter subscribers (active)", "pl_starter", [0, 2, 4, 6, 9, 12, 15, 18, 21, 24, 27, 30]),
    ("v_pro", "Pro subscribers (active)", "pl_pro", [0, 0, 1, 2, 3, 4, 5, 6, 8, 9, 11, 12]),
    ("v_business", "Business subscribers (active)", "pl_business", [0, 0, 0, 0, 1, 1, 1, 2, 2, 2, 3, 3]),
]
TEAM = [
    ("t_sales", "Sales / BD executive (+ commission above)", [0, 35000] + [35000] * 10),
    ("t_sales2", "Second sales executive", [0] * 5 + [35000] * 7),
    ("t_eng", "Full-stack / 3D engineer (after grant or seed)", [0] * 4 + [100000] * 8),
    ("t_ops", "Ops lead (QA, onboarding, support)", [0] * 3 + [30000] * 9),
]

# ---------------------------------------------------------------- cell registry + evaluator
refs = {}      # key -> (sheet, col_letter, row)
values = {}    # (sheet, "C5") -> value or formula
order = []


def put(sheet, col, row, v, key=None):
    values[(sheet, f"{col}{row}")] = v
    if key:
        refs[key] = (sheet, col, row)


def resolve(formula, sheet, col):
    def sub(m):
        key, rel = m.group(1), m.group(2)
        s, c, r = refs[key]
        c = col if rel else c
        cell = f"{c}{r}" if rel else f"${c}${r}"
        return cell if s == sheet else f"{s}!{cell}"
    return re.sub(r"\{(\w+)(:col)?\}", sub, formula)


cache = {}
REF = re.compile(r"(?:(\w+)!)?\$?([A-Z]{1,2})\$?(\d+)(?::\$?([A-Z]{1,2})\$?(\d+))?")


def col_range(a, b):
    from openpyxl.utils import column_index_from_string as ci
    return [get_column_letter(i) for i in range(ci(a), ci(b) + 1)]


def ev(sheet, cell):
    k = (sheet, cell)
    if k in cache:
        return cache[k]
    v = values.get(k, 0)
    if isinstance(v, str) and v.startswith("="):
        expr = v[1:]

        def rep(m):
            s = m.group(1) or sheet
            if m.group(4):
                cells = [f"{c}{r}" for c in col_range(m.group(2), m.group(4)) for r in range(int(m.group(3)), int(m.group(5)) + 1)]
                return "[" + ",".join(f"_ev({s!r},{c!r})" for c in cells) + "]"
            return f"_ev({s!r},'{m.group(2)}{m.group(3)}')"
        py = REF.sub(rep, expr).replace("^", "**")
        env = {"_ev": ev, "SUM": lambda *a: sum(sum(x) if isinstance(x, list) else x for x in a),
               "MIN": lambda *a: min(y for x in a for y in (x if isinstance(x, list) else [x])),
               "MAX": lambda *a: max(y for x in a for y in (x if isinstance(x, list) else [x])),
               "ROUNDUP": lambda x, n=0: math.ceil(x * 10 ** n) / 10 ** n,
               "ROUND": lambda x, n=0: round(x, n)}
        v = eval(py, env)  # noqa: S307 (our own formulas)
    elif v is None or isinstance(v, str):
        v = v
    cache[k] = v
    return v


# ---------------------------------------------------------------- layout
sheets = {}


def table(name, header, rows, widths):
    sheets[name] = (header, rows, widths)


# Inputs (values in column C)
rows = []
r = 2
for key, label, val, unit, src in INPUTS:
    if label is None:
        rows.append((r, [key], "section"))
    else:
        put("Inputs", "C", r, val, key)
        rows.append((r, [label, None, unit, src], "input"))
    r += 1
table("Inputs", ["Input", "", "Value", "Unit", "Source / confidence"], rows, [52, 2, 12, 20, 70])

# UnitCost (formulas in column C)
rows, r = [], 2
for key, label, f, unit, how in UNIT:
    put("UnitCost", "C", r, resolve(f, "UnitCost", "C"), key)
    rows.append((r, [label, None, unit, how], "formula"))
    r += 1
table("UnitCost", ["Cost item", "", "₹", "Unit", "How"], rows, [52, 2, 12, 22, 40])

# PriceList: B unit, C price, D direct cost, E pay fee, F commission, G contribution, H margin, I incl GST, J ₹/sqft, K notes
rows, r = [], 2
for key, product, unit, price, costf, comm, sqft, notes in PRICES:
    put("PriceList", "C", r, price, key)
    put("PriceList", "D", r, resolve(costf, "PriceList", "D"))
    put("PriceList", "E", r, f"=C{r}*{resolve('{pg_fee}', 'PriceList', 'E')}")
    put("PriceList", "F", r, f"=C{r}*{resolve('{sales_comm}', 'PriceList', 'F')}" if comm else 0)
    put("PriceList", "G", r, f"=C{r}-D{r}-E{r}-F{r}", key + "_contrib")
    put("PriceList", "H", r, f"=G{r}/C{r}")
    put("PriceList", "I", r, f"=C{r}*(1+{resolve('{gst}', 'PriceList', 'I')})")
    put("PriceList", "J", r, f"=C{r}/{resolve('{sqft_home}', 'PriceList', 'J')}" if sqft else "")
    rows.append((r, [product, unit], "price", notes))
    r += 1
table("PriceList", ["Product", "Unit", "Price ₹ (ex-GST)", "Direct cost ₹", "Payment fee ₹", "Sales commission ₹",
                    "Contribution ₹", "Margin %", "Price incl. 18% GST ₹", "₹ per sq ft (2BHK)", "Notes"],
      rows, [70, 18, 14, 13, 12, 14, 14, 10, 16, 14, 44])

# Plans: B ₹/mo, C ₹/yr, D tours, E seats, F walks, G hosting, H re-capture, I OTP, J walk cost, K support, L pay fee, M COGS, N GM, O GM%, P features, Q anchor
rows, r = [], 2
for key, plan, price, tours, seats, walks, sup, feats, anchor in PLANS:
    put("Plans", "B", r, price, key)
    put("Plans", "C", r, f"=B{r}*10")
    put("Plans", "D", r, tours)
    put("Plans", "E", r, seats)
    put("Plans", "F", r, walks)
    put("Plans", "G", r, f"=D{r}*{resolve('{host_tour_m}', 'Plans', 'G')}")
    put("Plans", "H", r, f"=D{r}*{resolve('{tour_churn}*{rooms_home}*({cpu_room}+{store_raw_room})', 'Plans', 'H')}")
    put("Plans", "I", r, f"=E{r}*{resolve('{otp_cust_m}', 'Plans', 'I')}")
    put("Plans", "J", r, f"=F{r}*{resolve('{cost_walk}', 'Plans', 'J')}")
    put("Plans", "K", r, "=" + resolve("{" + sup + "}", "Plans", "K"))
    put("Plans", "L", r, f"=B{r}*{resolve('{pg_fee}', 'Plans', 'L')}")
    put("Plans", "M", r, f"=SUM(G{r}:L{r})")
    put("Plans", "N", r, f"=B{r}-M{r}", key + "_contrib")
    put("Plans", "O", r, f"=N{r}/B{r}")
    rows.append((r, [plan], "plan", feats, anchor))
    r += 1
table("Plans", ["Plan", "₹ / month (ex-GST)", "₹ / year (2 months free)", "Active tours", "Seats", "3D walkthroughs / month",
                "Hosting ₹", "Re-capture processing ₹", "Login SMS ₹", "3D GPU ₹", "Support ₹", "Payment fee ₹",
                "Cost / month ₹", "Gross margin ₹", "Gross margin %", "Features", "Competitor anchor"],
      rows, [18, 14, 16, 10, 8, 12, 10, 12, 10, 10, 10, 11, 12, 13, 11, 50, 52])

# Competitors
rows, r = [], 2
for name, plan, lo, hi, what, src, conf in COMPETITORS:
    put("Competitors", "C", r, lo)
    put("Competitors", "D", r, hi)
    put("Competitors", "E", r, f"=C{r}*{resolve('{fx}', 'Competitors', 'E')}")
    put("Competitors", "F", r, f"=D{r}*{resolve('{fx}', 'Competitors', 'F')}")
    rows.append((r, [name, plan], "comp", what, src, conf))
    r += 1
r += 1
other_start = r
for name, unit, price, src, conf in COMP_OTHER:
    rows.append((r, [name, unit], "comp_other", price, src, conf))
    r += 1
table("Competitors", ["Competitor", "Plan / unit", "US$ / month (low)", "US$ / month (high)", "₹ / month (low)", "₹ / month (high)",
                      "What you get", "Source", "Confidence"], rows, [24, 16, 14, 14, 14, 14, 44, 60, 50])

# Plan12M: A label, B unit price, C unit contribution, D..O months, P total
mcols = [get_column_letter(4 + i) for i in range(MONTHS)]
rows, r = [], 2
rows.append((r, ["VOLUMES (targets: edit the blue numbers)"], "section")); r += 1
vol_rows = {}
for key, label, pk, vols in VOLUMES:
    sheet = "Plans" if pk.startswith("pl_") else "PriceList"
    put("Plan12M", "B", r, "=" + resolve("{" + pk + "}", "Plan12M", "B"))
    put("Plan12M", "C", r, "=" + resolve("{" + pk + "_contrib}", "Plan12M", "C"))
    for c, v in zip(mcols, vols):
        put("Plan12M", c, r, v)
    put("Plan12M", "P", r, f"=SUM(D{r}:O{r})")
    refs[key] = ("Plan12M", "D", r)
    vol_rows[key] = r
    rows.append((r, [label], "volume"))
    r += 1
r += 1
rev_r = r
for c in mcols + ["P"]:
    put("Plan12M", c, r, "=(" + "+".join(f"{c}{vr}*$B${vr}" for vr in vol_rows.values()) + ")*" + resolve("{vol_scale}", "Plan12M", c) if c != "P" else f"=SUM(D{r}:O{r})")
rows.append((r, ["Revenue ₹ (ex-GST)"], "total")); r += 1
con_r = r
for c in mcols + ["P"]:
    put("Plan12M", c, r, "=(" + "+".join(f"{c}{vr}*$C${vr}" for vr in vol_rows.values()) + ")*" + resolve("{vol_scale}", "Plan12M", c) if c != "P" else f"=SUM(D{r}:O{r})")
rows.append((r, ["Contribution ₹ (after direct costs, fees, commission)"], "total")); r += 1
r += 1
rows.append((r, ["FIXED COSTS ₹ / month"], "section")); r += 1
fixed_start = r
for key, label in [("fx_infra", "Cloud infrastructure"), ("fx_tools", "Tools"), ("fx_ca", "Accounting / compliance"), ("fx_marketing", "Marketing")]:
    for c in mcols:
        put("Plan12M", c, r, "=" + resolve("{" + key + "}", "Plan12M", c))
    put("Plan12M", "P", r, f"=SUM(D{r}:O{r})")
    rows.append((r, [label], "fixed")); r += 1
for key, label, vals in TEAM:
    for c, v in zip(mcols, vals):
        put("Plan12M", c, r, v)
    put("Plan12M", "P", r, f"=SUM(D{r}:O{r})")
    rows.append((r, [label], "volume")); r += 1
fixed_end = r - 1
fix_r = r
for c in mcols + ["P"]:
    put("Plan12M", c, r, f"=SUM({c}{fixed_start}:{c}{fixed_end})")
rows.append((r, ["Total fixed costs ₹"], "total")); r += 1
r += 1
ebit_r = r
for c in mcols + ["P"]:
    put("Plan12M", c, r, f"={c}{con_r}-{c}{fix_r}")
rows.append((r, ["Monthly profit / (loss) ₹"], "total")); r += 1
cash_r = r
for i, c in enumerate(mcols):
    prev = resolve("{start_cash}", "Plan12M", c) if i == 0 else f"{mcols[i-1]}{r}"
    put("Plan12M", c, r, f"={prev}+{c}{ebit_r}")
rows.append((r, ["Cash in bank ₹ (start + cumulative profit)"], "total")); r += 1
low_r = r
put("Plan12M", "D", r, f"=MIN(D{cash_r}:O{cash_r})")
rows.append((r, ["Lowest cash point ₹ (negative = money you must raise first)"], "total")); r += 1
cap_r = r
for c in mcols:
    put("Plan12M", c, r, f"=ROUNDUP(({c}{vol_rows['v_dfy']}+{c}{vol_rows['v_villa']}*2)/({resolve('{homes_per_day}*{days_month}', 'Plan12M', c)}),0)")
rows.append((r, ["Field capturers needed (cost already in done-for-you price)"], "volume_calc")); r += 1
table("Plan12M", ["Line", "Unit price ₹", "Contribution / unit ₹"] + [f"M{i+1}" for i in range(MONTHS)] + ["12-month total"],
      rows, [56, 13, 16] + [11] * MONTHS + [15])

# BreakEven (column C)
rows, r = [], 2
be = [
    ("be_fixed_lean", "Fixed costs, lean (no salaries)", "={fx_infra}+{fx_tools}+{fx_ca}+{fx_marketing}", "₹ / month"),
    ("be_fixed_team", "Fixed costs with 2 sales + engineer + ops lead", "={be_fixed_lean}+35000*2+100000+30000", "₹ / month"),
    ("be_dfy_lean", "Done-for-you tours/month to cover lean costs", "={be_fixed_lean}/{p_dfy_contrib}", "tours"),
    ("be_dfy_team", "Done-for-you tours/month to cover full team", "={be_fixed_team}/{p_dfy_contrib}", "tours"),
    ("be_pro_team", "…or Pro subscribers to cover full team", "={be_fixed_team}/{pl_pro_contrib}", "subscribers"),
    ("be_dev_team", "…or developer packages per month to cover full team", "={be_fixed_team}/{p_dev_contrib}", "projects"),
    ("be_dfy_capacity", "Homes one capturer can shoot per month", "={homes_per_day}*{days_month}", "homes"),
]
for key, label, f, unit in be:
    put("BreakEven", "C", r, resolve(f, "BreakEven", "C"), key)
    rows.append((r, [label, None, unit], "formula")); r += 1
table("BreakEven", ["Question", "", "Answer", "Unit"], rows, [60, 2, 14, 14])

# ---------------------------------------------------------------- write xlsx + csv
wb = Workbook()
wb.remove(wb.active)
readme = wb.create_sheet("README")
for i, line in enumerate([
    "Nazariya: unit pricing and unit economics (all prices in ₹, excluding 18% GST)",
    "",
    "How to use: change any BLUE number on 'Inputs' or 'Plan12M'. Everything else is a formula and recalculates.",
    "Sheets: Inputs → UnitCost (what one unit costs us) → PriceList (what we charge, margin) → Plans (SaaS tiers) →",
    "Competitors (public price points with sources) → Plan12M (12-month revenue, costs, cash) → BreakEven.",
    "",
    "Measured on this app: stitching time per room (14.9 s), raw capture size (3.9 MB / room), published size (~1 MB / room).",
    "NOT measured yet: GPU minutes per 3D walkthrough (recon scripts have never run on a GPU), real visitor numbers,",
    "capturer productivity. Treat those as assumptions until the first 20 paid jobs.",
    "Competitor prices are public or third-party figures and often conflict; see the Confidence column.",
    "Built 2026-10-17 by tools/build_pricing.py.",
], start=1):
    readme.cell(row=i, column=1, value=line).font = BOLD if i == 1 else Font()
readme.column_dimensions["A"].width = 120

for name, (header, rows, widths) in sheets.items():
    ws = wb.create_sheet(name)
    for j, h in enumerate(header, start=1):
        c = ws.cell(row=1, column=j, value=h)
        c.font, c.fill = BOLD, HEAD
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "B2"
    for row, texts, kind, *extra in rows:
        for j, t in enumerate(texts, start=1):
            if t is not None:
                ws.cell(row=row, column=j, value=t)
        if kind in ("section", "total"):
            ws.cell(row=row, column=1).font = BOLD
        if name == "PriceList":
            ws.cell(row=row, column=11, value=extra[0])
        if name == "Plans":
            ws.cell(row=row, column=16, value=extra[0])
            ws.cell(row=row, column=17, value=extra[1])
        if name == "Competitors":
            if kind == "comp":
                ws.cell(row=row, column=7, value=extra[0])
                ws.cell(row=row, column=8, value=extra[1])
                ws.cell(row=row, column=9, value=extra[2])
            else:
                ws.cell(row=row, column=7, value=extra[0])
                ws.cell(row=row, column=8, value=extra[1])
                ws.cell(row=row, column=9, value=extra[2])
    for (s, cell), v in values.items():
        if s != name:
            continue
        c = ws[cell]
        c.value = v
        typed = not (isinstance(v, str) and v.startswith("="))
        if typed and v != "":
            c.font = BLUE
        col = re.match(r"[A-Z]+", cell).group(0)
        if (name, col) in {("PriceList", "H"), ("Plans", "O")} or (name == "Inputs" and isinstance(v, float) and v < 1 and refs and any(
                refs[k] == ("Inputs", "C", int(cell[1:])) and k in ("gst", "gpu_retry", "pg_fee", "sales_comm", "tour_churn") for k in refs)):
            c.number_format = "0%"
        elif not typed or isinstance(v, (int, float)):
            c.number_format = "#,##0.00" if name == "UnitCost" or (name == "Inputs" and isinstance(v, float)) else "#,##0"
wb.save(OUT / "nazariya-unit-economics.xlsx")

# CSV snapshot (computed values)
def fmt(v):
    if isinstance(v, float):
        return f"{v:.2f}" if abs(v) < 100 else f"{v:.0f}"
    return "" if v is None else v

for name, (header, rows, _w) in sheets.items():
    with open(OUT / "csv" / f"{name}.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow([h for h in header])
        ncol = len(header)
        for row, texts, kind, *extra in rows:
            line = [""] * ncol
            for j, t in enumerate(texts):
                if t is not None:
                    line[j] = t
            for j in range(ncol):
                cell = f"{get_column_letter(j + 1)}{row}"
                if (name, cell) in values and values[(name, cell)] != "":
                    line[j] = fmt(ev(name, cell))
            if name == "PriceList":
                line[10] = extra[0]
            if name == "Plans":
                line[15], line[16] = extra
            if name == "Competitors":
                line[6], line[7], line[8] = extra
            w.writerow(line)

# Print a summary for the console.
def show(sheet, key=None, cell=None):
    s, c, r = refs[key] if key else (sheet, cell[0], int(cell[1:]))
    return ev(s, f"{c}{r}")

print("Unit costs (₹):")
for key, label, *_ in UNIT:
    print(f"  {label:58s} {show(None, key):10.2f}")
print("Price list (₹ ex-GST): price | direct cost | contribution | margin")
for key, product, *_ in PRICES:
    r = refs[key][2]
    print(f"  {product[:60]:60s} {ev('PriceList', f'C{r}'):8.0f} {ev('PriceList', f'D{r}'):8.0f} {ev('PriceList', f'G{r}'):8.0f} {ev('PriceList', f'H{r}'):6.0%}")
print("Plans: price | cost | GM%")
for key, plan, *_ in PLANS:
    r = refs[key][2]
    print(f"  {plan:20s} {ev('Plans', f'B{r}'):8.0f} {ev('Plans', f'M{r}'):8.0f} {ev('Plans', f'O{r}'):6.0%}")
print("12 months: revenue", round(ev("Plan12M", f"P{rev_r}")), "contribution", round(ev("Plan12M", f"P{con_r}")),
      "fixed", round(ev("Plan12M", f"P{fix_r}")), "profit", round(ev("Plan12M", f"P{ebit_r}")))
print("  monthly revenue:", [round(ev("Plan12M", f"{c}{rev_r}")) for c in mcols])
print("  monthly profit: ", [round(ev("Plan12M", f"{c}{ebit_r}")) for c in mcols])
print("  cash:           ", [round(ev("Plan12M", f"{c}{cash_r}")) for c in mcols])
print("  lowest cash:", round(ev("Plan12M", f"D{low_r}")), " capturers:", [ev("Plan12M", f"{c}{cap_r}") for c in mcols])
for key, label, *_ in be:
    print(f"  {label:58s} {show(None, key):10.1f}")
